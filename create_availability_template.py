"""Generate a generic NOC availability workbook for Google Sheets.

Output (next to this script):
  templates/NOC_Availability_Template.xlsx
  templates/NOC_Availability_Template.csv

The first tab is named זמינות so Google Sheets gid=0 matches what the app imports.
"""
from __future__ import annotations

import csv
import os
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SHIFT_TYPES = ["בוקר", "ערב", "לילה", "כונן"]
DAY_NAMES = ["ראשון", "שני", "שלישי", "רביעי", "חמישי", "שישי", "שבת"]
EMPLOYEES = [f"עובד {i}" for i in range(1, 8)]
CHECK = "✔️"
STAR = "⭐"
NOTE = "שמירה על הילדים"

# Per employee: 7 days × 4 shifts. Values: "", CHECK, STAR, or NOTE.
# Dense enough for auto-schedule to fill a week; nights are spread (not consecutive).
SAMPLE = {
    "עובד 1": [
        [CHECK, CHECK, STAR, ""],
        [CHECK, CHECK, "", ""],
        [CHECK, "", "", CHECK],
        [CHECK, CHECK, CHECK, ""],
        [CHECK, CHECK, "", ""],
        [CHECK, STAR, "", ""],
        [CHECK, "", "", CHECK],
    ],
    "עובד 2": [
        ["", CHECK, "", CHECK],
        [CHECK, CHECK, CHECK, ""],
        [CHECK, CHECK, "", ""],
        ["", CHECK, "", CHECK],
        [CHECK, CHECK, "", ""],
        ["", CHECK, CHECK, ""],
        ["", CHECK, "", STAR],
    ],
    "עובד 3": [
        [CHECK, "", "", ""],
        [CHECK, CHECK, STAR, ""],
        [CHECK, CHECK, "", CHECK],
        [CHECK, "", CHECK, ""],
        [CHECK, CHECK, "", ""],
        [CHECK, "", "", CHECK],
        [CHECK, CHECK, "", ""],
    ],
    "עובד 4": [
        [CHECK, CHECK, "", ""],
        [CHECK, STAR, "", CHECK],
        ["", CHECK, CHECK, ""],
        [CHECK, CHECK, "", ""],
        [CHECK, "", STAR, CHECK],
        [CHECK, CHECK, "", ""],
        ["", CHECK, CHECK, ""],
    ],
    "עובד 5": [
        [NOTE, CHECK, CHECK, ""],
        ["", CHECK, "", CHECK],
        [CHECK, CHECK, STAR, ""],
        ["", CHECK, "", ""],
        [CHECK, CHECK, CHECK, ""],
        ["", "", "", CHECK],
        [CHECK, CHECK, STAR, ""],
    ],
    "עובד 6": [
        [CHECK, "", "", CHECK],
        [CHECK, CHECK, "", ""],
        [CHECK, "", "", ""],
        [CHECK, CHECK, "", STAR],
        [CHECK, "", "", CHECK],
        [STAR, CHECK, "", ""],
        [CHECK, "", CHECK, ""],
    ],
    "עובד 7": [
        ["", STAR, "", ""],
        ["", CHECK, "", ""],
        ["", CHECK, CHECK, ""],
        ["", CHECK, "", ""],
        ["", CHECK, "", CHECK],
        ["", CHECK, CHECK, ""],
        ["", STAR, "", CHECK],
    ],
}


def _upcoming_sunday(today: date | None = None) -> date:
    today = today or date.today()
    # Monday=0 ... Sunday=6; stay on today if it is already Sunday.
    return today + timedelta(days=(6 - today.weekday()) % 7)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _thin():
    return Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )


def build_workbook(sunday: date) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "זמינות"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B4"

    header_fill = PatternFill("solid", fgColor="1A237E")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    shift_fill = PatternFill("solid", fgColor="3949AB")
    name_fill = PatternFill("solid", fgColor="EEF2FF")
    check_fill = PatternFill("solid", fgColor="C8E6C9")
    star_fill = PatternFill("solid", fgColor="FFF9C4")
    note_fill = PatternFill("solid", fgColor="FFE0B2")
    empty_fill = PatternFill("solid", fgColor="FFFFFF")

    ws["A1"] = ""
    ws["A2"] = "יום"
    ws["A3"] = "משמרת"
    for cell in (ws["A1"], ws["A2"], ws["A3"]):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = _center()
        cell.border = _thin()
    ws.column_dimensions["A"].width = 16

    for d, day_name in enumerate(DAY_NAMES):
        start = 2 + d * 4
        end = start + 3
        day_date = sunday + timedelta(days=d)
        date_text = day_date.strftime("%d/%m/%Y")

        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)

        date_cell = ws.cell(1, start, date_text)
        date_cell.font = header_font
        date_cell.fill = header_fill
        date_cell.alignment = _center()

        day_cell = ws.cell(2, start, day_name)
        day_cell.font = header_font
        day_cell.fill = header_fill
        day_cell.alignment = _center()

        for offset, shift in enumerate(SHIFT_TYPES):
            col = start + offset
            cell = ws.cell(3, col, shift)
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
            cell.fill = shift_fill
            cell.alignment = _center()
            cell.border = _thin()
            ws.column_dimensions[get_column_letter(col)].width = 12
            for r in (1, 2):
                ws.cell(r, col).border = _thin()
                ws.cell(r, col).fill = header_fill

    ws["A1"].comment = Comment(
        "תבנית זמינות ל-NOC Shift Scheduler.\n"
        "שנה שמות עובדים ותאריכים, סמן ✔️ / ⭐, שתף כ-Anyone with the link, "
        "והדבק את כתובת הגיליון באפליקציה.",
        "NOC Shift Scheduler",
    )

    for row_idx, name in enumerate(EMPLOYEES, start=4):
        name_cell = ws.cell(row_idx, 1, name)
        name_cell.font = Font(name="Calibri", bold=True, size=12, color="1E293B")
        name_cell.fill = name_fill
        name_cell.alignment = Alignment(horizontal="right", vertical="center")
        name_cell.border = _thin()
        for d in range(7):
            for s_idx, marker in enumerate(SAMPLE[name][d]):
                cell = ws.cell(row_idx, 2 + d * 4 + s_idx, marker)
                cell.alignment = _center()
                cell.font = Font(name="Calibri", size=14)
                cell.border = _thin()
                if marker == STAR:
                    cell.fill = star_fill
                elif marker == CHECK:
                    cell.fill = check_fill
                elif marker:
                    cell.fill = note_fill
                    cell.font = Font(name="Calibri", size=9, color="E65100")
                else:
                    cell.fill = empty_fill

    last_data_col = 1 + 7 * 4
    last_data_row = 3 + len(EMPLOYEES)
    dv = DataValidation(
        type="list",
        formula1='"✔️,⭐,V,1,כן"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=False,
    )
    dv.prompt = "✔️ זמין · ⭐ העדפה · ריק = לא זמין · טקסט חופשי = הערה"
    dv.promptTitle = "סימון זמינות"
    ws.add_data_validation(dv)
    dv.add(f"B4:{get_column_letter(last_data_col)}{last_data_row + 8}")

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 20
    for r in range(4, last_data_row + 1):
        ws.row_dimensions[r].height = 26

    # Extra blank employee rows so a new team can add people immediately.
    for extra in range(1, 4):
        r = last_data_row + extra
        cell = ws.cell(r, 1, "")
        cell.fill = name_fill
        cell.border = _thin()
        for c in range(2, last_data_col + 1):
            empty = ws.cell(r, c, "")
            empty.fill = empty_fill
            empty.border = _thin()
            empty.alignment = _center()

    instr = wb.create_sheet("הוראות")
    instr.sheet_view.rightToLeft = True
    instr.column_dimensions["A"].width = 88
    instr["A1"] = "איך להשתמש בתבנית הזו תוך דקה"
    instr["A1"].font = Font(name="Calibri", bold=True, size=18, color="1A237E")
    steps = [
        "",
        "1. ב-Google Sheets: קובץ → ייבוא → העלאה → בחר את NOC_Availability_Template.xlsx",
        "   (או פתח ב-Excel, סמן את גיליון «זמינות», והעתק לגיליון חדש).",
        "2. השאר את הלשונית «זמינות» ראשונה. האפליקציה קוראת את הלשונית הראשונה (gid=0).",
        "3. שנה את השמות «עובד 1»…«עובד 7» לשמות אמיתיים. מחק או הוסף שורות לפי גודל הצוות.",
        "4. עדכן את התאריכים בשורה 1. אל תשנה את כותרות המשמרות בשורה 3 (בוקר / ערב / לילה / כונן).",
        "5. סמנים:",
        "     ✔️   או  V / 1 / כן     = זמין למשמרת",
        "     ⭐   או  P / עדיפות      = זמין + עדיפות (בונוס בשיבוץ האוטומטי)",
        "     ריק                     = לא זמין",
        "     כל טקסט אחר             = הערה בלבד, לא נספר כזמינות",
        "6. שיתוף: שתף → גישה כללית → כל מי שיש לו את הקישור → צופה.",
        "   בלי זה האפליקציה תקבל דף התחברות במקום CSV.",
        "7. העתק את כתובת הגיליון לאפליקציית NOC Shift Scheduler ולחץ «ייבוא נתונים».",
        "",
        "הדוגמה כבר מסומנת כך שאפשר לייבא וללחוץ «שבץ אוטומטי» מיד, בלי למלא כלום.",
        "אחרי זה תחליף שמות ותסמן מחדש לפי הזמינות האמיתית של השבוע.",
    ]
    for i, line in enumerate(steps, start=2):
        instr.cell(i, 1, line).font = Font(name="Calibri", size=13, color="1E293B")
        instr.row_dimensions[i].height = 22

    return wb


def write_csv(path: str, sunday: date) -> None:
    rows: list[list[str]] = []
    dates = [""]
    days = ["יום"]
    shifts = ["משמרת"]
    for d, day_name in enumerate(DAY_NAMES):
        date_text = (sunday + timedelta(days=d)).strftime("%d/%m/%Y")
        dates.extend([date_text, "", "", ""])
        days.extend([day_name, "", "", ""])
        shifts.extend(SHIFT_TYPES)
    rows.append(dates)
    rows.append(days)
    rows.append(shifts)
    for name in EMPLOYEES:
        row = [name]
        for d in range(7):
            row.extend(SAMPLE[name][d])
        rows.append(row)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)


def main() -> None:
    here = os.path.dirname(os.path.abspath(__file__))
    out_dir = os.path.join(here, "templates")
    os.makedirs(out_dir, exist_ok=True)
    sunday = _upcoming_sunday()
    xlsx_path = os.path.join(out_dir, "NOC_Availability_Template.xlsx")
    csv_path = os.path.join(out_dir, "NOC_Availability_Template.csv")
    wb = build_workbook(sunday)
    wb.save(xlsx_path)
    write_csv(csv_path, sunday)
    print(f"Wrote {xlsx_path}")
    print(f"Wrote {csv_path}")
    print(f"Week starting Sunday {sunday.strftime('%d/%m/%Y')}")


if __name__ == "__main__":
    main()
