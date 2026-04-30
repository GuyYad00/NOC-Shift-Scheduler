# pyre-ignore-all-errors
"""NOC Shift Scheduler - Premium GUI with RTL fixes and Heatmap."""

import threading
import webbrowser
import urllib.parse
import customtkinter as ctk
from tkinter import messagebox, filedialog
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from typing import Optional, Dict, Any

from config import (
    APP_TITLE, APP_WIDTH, APP_HEIGHT,
    SHIFT_TYPES,
    WHATSAPP_MSG_HEADER, COLOR_HEATMAP_EMPTY, COLOR_HEATMAP_LOW,
    slots_for_day_shift, normalize_assignment_key,
)
from sheets_reader import fetch_and_parse
from data_store import (
    load_settings, save_settings,
    get_current_month_scores, get_current_month_split, update_month_scores_split,
    update_month_scores, reset_month_scores,
    reset_all_historical_data,
    get_prev_week_data, save_last_week_data, get_historical_stats
)
from scheduler import schedule_shifts, calculate_split_week_scores

# Appearance settings
ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

# Professional Color Palette
EMPTY_LABEL = "-- ריק --"
COLOR_MANUAL = "#42A5F5"      # Blue
COLOR_AUTO = "#66BB6A"        # Green
COLOR_UNFILLED = "#EF5350"    # Red
COLOR_WARNING = "#FFA726"     # Orange
COLOR_WHATSAPP = "#25D366"    # WhatsApp
COLOR_CARD_BG = "#F5F5F7"     # Light Apple-style Gray for cards

# שורות תצוגה: (תווית, סוג משמרת, מספר סלוט)
SCHEDULE_ROW_DEFS = [
    ("בוקר (א)", "בוקר", 0),
    ("בוקר (ב)", "בוקר", 1),
    ("ערב", "ערב", 0),
    ("לילה", "לילה", 0),
    ("כונן", "כונן", 0),
]


class StatsWindow(ctk.CTkToplevel):
    """Premium analytics window for historical data."""
    def __init__(self, parent):
        super().__init__(master=parent) # type: ignore
        self.title("לוח בקרה וסטטיסטיקה")
        self.geometry("600x580")
        self.grab_set()
        self.configure(fg_color="#FAFBFC")

        # Title bar
        title = ctk.CTkFrame(self, fg_color="#1565C0", corner_radius=0, height=50)
        title.pack(fill="x")
        title.pack_propagate(False)
        ctk.CTkLabel(title, text="📊  סיכום ניקוד היסטורי", font=("Calibri", 19, "bold"), text_color="white").pack(expand=True)
        
        scroll = ctk.CTkScrollableFrame(self, width=550, height=460, fg_color="transparent")
        scroll.pack(padx=15, pady=15, fill="both", expand=True)

        all_stats = get_historical_stats()
        if not all_stats:
            ctk.CTkLabel(scroll, text="אין עדיין נתונים היסטוריים שמורים", font=("Calibri", 14), text_color="#94A3B8").pack(pady=50)
            return

        for month, scores in sorted(all_stats.items(), reverse=True):
            frame = ctk.CTkFrame(scroll, corner_radius=12, fg_color="#F0F4F8", border_width=1, border_color="#E0E4EA")
            frame.pack(fill="x", padx=8, pady=8)
            
            # Month header
            mhdr = ctk.CTkFrame(frame, fg_color="transparent")
            mhdr.pack(fill="x", padx=15, pady=(10, 5))
            ctk.CTkLabel(mhdr, text=f"📅  חודש: {month}", font=("Calibri", 14, "bold"), text_color="#1565C0", anchor="e").pack(side="right")
            
            ctk.CTkFrame(frame, fg_color="#1565C0", height=2, corner_radius=1).pack(fill="x", padx=15, pady=2)

            for emp, score in sorted(scores.items(), key=lambda x: -x[1]):
                line = ctk.CTkFrame(frame, fg_color="transparent")
                line.pack(fill="x", padx=18, pady=3)
                ctk.CTkLabel(line, text=f"{score:.1f} נק'", font=("Calibri", 12), text_color="#475569").pack(side="left")
                ctk.CTkLabel(line, text=emp, font=("Calibri", 12, "bold"), text_color="#1E293B", width=150, anchor="e").pack(side="right")


class SettingsDialog(ctk.CTkToplevel):
    """Premium settings dialog with card-based sections."""
    
    # Section colors
    _CARD_BG = "#F0F4F8"
    _ACCENT_BLUE = "#1565C0"
    _ACCENT_TEAL = "#00897B"
    _ACCENT_ORANGE = "#EF6C00"
    _ACCENT_RED = "#C62828"
    
    def __init__(self, parent, employees: list, excluded: list, settings: dict):
        super().__init__(master=parent) # type: ignore
        self.title("הגדרות מערכת")
        self.geometry("500x780")
        self.resizable(False, False)
        self.grab_set()
        self.configure(fg_color="#FAFBFC")

        self.result_excluded = list(excluded)
        self.settings = settings
        self.check_vars = {}
        self.parent_app = parent

        # --- Title ---
        title_frame = ctk.CTkFrame(self, fg_color=self._ACCENT_BLUE, corner_radius=0, height=55)
        title_frame.pack(fill="x")
        title_frame.pack_propagate(False)
        ctk.CTkLabel(title_frame, text="⚙️  הגדרות מערכת", font=("Calibri", 20, "bold"), text_color="white").pack(expand=True)

        # Scrollable body
        body = ctk.CTkScrollableFrame(self, fg_color="transparent", corner_radius=0)
        body.pack(fill="both", expand=True, padx=0, pady=0)

        # ═══════════════════ Card 1: Algorithm Constraints ═══════════════════
        card1 = self._make_card(body)
        self._section_header(card1, "📐  אילוצי אלגוריתם", self._ACCENT_BLUE)
        
        self._add_constraint_row(card1, "\u200Fמינימום משמרות לעובד\u200F:", "min_shifts", 3)
        self._add_constraint_row(card1, "\u200Fמקסימום ימי עבודה ברצף\u200F:", "max_consecutive_days", 6)

        # ═══════════════════ Card 2: Employee Exclusion ═══════════════════
        card2 = self._make_card(body)
        self._section_header(card2, "👤  החרגת עובדים מאיזון", self._ACCENT_TEAL)

        scroll = ctk.CTkScrollableFrame(card2, width=380, height=180, fg_color="white", corner_radius=10, border_width=1, border_color="#E0E0E0")
        scroll.pack(padx=15, pady=(5, 15), fill="both", expand=True)

        for emp in employees:
            var = ctk.BooleanVar(value=emp in excluded)
            self.check_vars[emp] = var
            cb_frame = ctk.CTkFrame(scroll, fg_color="transparent")
            cb_frame.pack(fill="x", padx=10, pady=4)
            ctk.CTkCheckBox(
                cb_frame, text=emp, variable=var, font=("Calibri", 13),
                checkbox_width=22, checkbox_height=22, corner_radius=6,
                fg_color=self._ACCENT_TEAL, hover_color="#00796B",
                border_width=2
            ).pack(anchor="e")

        # ═══════════════════ Card 3: Data Reset ═══════════════════
        card3 = self._make_card(body)
        self._section_header(card3, "🗑️  איפוס נתונים", self._ACCENT_RED)
        
        ctk.CTkLabel(card3, text="פעולות אלו אינן ניתנות לביטול", font=("Calibri", 11), text_color="#999").pack(pady=(0, 8))

        reset_frame = ctk.CTkFrame(card3, fg_color="transparent")
        reset_frame.pack(fill="x", padx=15, pady=(0, 15))
        ctk.CTkButton(
            reset_frame, text="🔄  איפוס ניקוד חודשי", width=195,
            height=36, corner_radius=10,
            fg_color=self._ACCENT_ORANGE, hover_color="#E65100",
            font=("Calibri", 12, "bold"), command=self._reset_monthly
        ).pack(side="right", padx=5)
        ctk.CTkButton(
            reset_frame, text="⚠️  איפוס כל הסטטיסטיקות", width=195,
            height=36, corner_radius=10,
            fg_color=self._ACCENT_RED, hover_color="#B71C1C",
            font=("Calibri", 12, "bold"), command=self._reset_all
        ).pack(side="right", padx=5)

        # ═══════════════════ Bottom Action Bar ═══════════════════
        bottom_bar = ctk.CTkFrame(self, fg_color="#F0F0F2", corner_radius=0, height=65)
        bottom_bar.pack(fill="x", side="bottom")
        bottom_bar.pack_propagate(False)
        
        ctk.CTkButton(
            bottom_bar, text="ביטול", width=110, height=38, corner_radius=10,
            fg_color="transparent", border_width=2, border_color="#B0BEC5",
            text_color="#546E7A", hover_color="#ECEFF1",
            font=("Calibri", 13), command=self.destroy
        ).pack(side="left", padx=20, pady=13)
        ctk.CTkButton(
            bottom_bar, text="✓  שמור הגדרות", width=150, height=38, corner_radius=10,
            fg_color=self._ACCENT_BLUE, hover_color="#0D47A1",
            font=("Calibri", 14, "bold"), command=self._save
        ).pack(side="left", padx=5, pady=13)

    # ── Helpers ──
    def _make_card(self, parent):
        card = ctk.CTkFrame(parent, fg_color=self._CARD_BG, corner_radius=14, border_width=1, border_color="#E4E8ED")
        card.pack(fill="x", padx=18, pady=(12, 0))
        return card

    def _section_header(self, parent, text, color):
        hdr = ctk.CTkFrame(parent, fg_color="transparent")
        hdr.pack(fill="x", padx=15, pady=(12, 8))
        ctk.CTkLabel(hdr, text=text, font=("Calibri", 16, "bold"), text_color=color, anchor="e").pack(side="right")
        line = ctk.CTkFrame(parent, height=2, fg_color=color, corner_radius=1)
        line.pack(fill="x", padx=15, pady=(0, 8))

    def _add_constraint_row(self, parent, label_text, key, default):
        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(fill="x", padx=20, pady=6)
        entry = ctk.CTkEntry(row, width=65, height=34, justify="center", corner_radius=8, border_width=1, border_color="#B0BEC5", font=("Calibri", 14))
        entry.insert(0, str(self.settings.get(key, default)))
        entry.pack(side="left")
        setattr(self, f"{key}_entry", entry)
        ctk.CTkLabel(row, text=label_text, font=("Calibri", 13), text_color="#37474F").pack(side="right", padx=5)

    def _reset_monthly(self):
        if messagebox.askyesno("אישור איפוס", "האם לאפס את הניקוד החודשי?\nפעולה זו אינה ניתנת לביטול.", parent=self):
            reset_month_scores()
            messagebox.showinfo("בוצע", "הניקוד החודשי אופס בהצלחה.", parent=self)

    def _reset_all(self):
        if messagebox.askyesno("אישור איפוס", "האם לאפס את כל הסטטיסטיקות ההיסטוריות?\nפעולה זו אינה ניתנת לביטול!", parent=self):
            reset_all_historical_data()
            messagebox.showinfo("בוצע", "כל הסטטיסטיקות אופסו בהצלחה.", parent=self)

    def _save(self):
        try:
            self.settings["min_shifts"] = int(self.min_shifts_entry.get())
            self.settings["max_consecutive_days"] = int(self.max_consecutive_days_entry.get())
            self.result_excluded = [e for e, v in self.check_vars.items() if v.get()]
            self.settings["excluded_employees"] = self.result_excluded
            self.destroy()
        except ValueError:
            messagebox.showerror("שגיאה", "נא להזין מספרים שלמים בלבד")


class App(ctk.CTk):
    """Main Application Controller — Premium UI."""
    
    # ── Design Tokens ──
    _BG = "#F5F7FA"
    _HEADER_BG = "#1A237E"
    _HEADER_ACCENT = "#3949AB"
    _CARD_BG = "#FFFFFF"
    _CARD_BORDER = "#E0E4EA"
    _TEXT_PRIMARY = "#1E293B"
    _TEXT_SECONDARY = "#64748B"
    _ACCENT_GREEN = "#059669"
    _ACCENT_BLUE = "#2563EB"
    _ACCENT_PURPLE = "#7C3AED"
    _BTN_RADIUS = 12
    _BTN_H = 38
    
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry(f"{APP_WIDTH}x{APP_HEIGHT + 40}")
        self.minsize(1200, 900)
        self.configure(fg_color=self._BG)

        self.week_data: Optional[Dict[str, Any]] = None
        self.manual_flags: Dict[tuple, bool] = {}
        self.current_assignments: Dict[tuple, str] = {}
        self.schedule_menus: Dict[tuple, Any] = {}
        self.settings: dict = load_settings()

        self._build_ui()
        self._load_saved_url()

    def _build_ui(self):
        # ═══════════════════ Top Header Bar ═══════════════════
        header = ctk.CTkFrame(self, fg_color=self._HEADER_BG, corner_radius=0, height=70)
        header.pack(fill="x")
        header.pack_propagate(False)
        
        # App title on right
        ctk.CTkLabel(
            header, text="📋  מערכת שיבוץ משמרות NOC",
            font=("Calibri", 18, "bold"), text_color="white"
        ).pack(side="right", padx=25)
        
        # URL input group on left
        url_group = ctk.CTkFrame(header, fg_color=self._HEADER_ACCENT, corner_radius=10, height=42)
        url_group.pack(side="left", padx=20, pady=14, fill="x", expand=True)
        url_group.pack_propagate(False)
        
        self.import_btn = ctk.CTkButton(
            url_group, text="📥  ייבוא נתונים", width=140, height=34,
            corner_radius=8, fg_color="#43A047", hover_color="#388E3C",
            font=("Calibri", 13, "bold"), command=self._import_clicked
        )
        self.import_btn.pack(side="left", padx=4, pady=4)
        
        self.url_entry = ctk.CTkEntry(
            url_group, placeholder_text="🔗 הדבק כאן את כתובת Google Sheets...",
            height=34, corner_radius=8, border_width=0,
            font=("Calibri", 12), fg_color="white", text_color=self._TEXT_PRIMARY
        )
        self.url_entry.pack(side="left", fill="x", expand=True, padx=(4, 8), pady=4)

        # ═══════════════════ Status Bar ═══════════════════
        status_bar = ctk.CTkFrame(self, fg_color="#EEF2FF", corner_radius=0, height=32)
        status_bar.pack(fill="x")
        status_bar.pack_propagate(False)
        self.status_label = ctk.CTkLabel(
            status_bar, text="● מוכן לעבודה", font=("Calibri", 12),
            text_color=self._TEXT_SECONDARY
        )
        self.status_label.pack(expand=True)

        # ═══════════════════ Main Table Card ═══════════════════
        self.table_outer = ctk.CTkFrame(
            self, corner_radius=16, fg_color=self._CARD_BG,
            border_width=1, border_color=self._CARD_BORDER
        )
        self.table_outer.pack(fill="both", expand=True, padx=20, pady=(10, 5))
        
        self.table_frame = ctk.CTkFrame(self.table_outer, fg_color="transparent")
        self.table_frame.pack(anchor="center", pady=15, padx=15)

        # ═══════════════════ Action Buttons Bar ═══════════════════
        actions = ctk.CTkFrame(self, fg_color=self._CARD_BG, corner_radius=14, border_width=1, border_color=self._CARD_BORDER)
        actions.pack(fill="x", padx=20, pady=8)
        actions_inner = ctk.CTkFrame(actions, fg_color="transparent")
        actions_inner.pack(fill="x", padx=15, pady=10)
        
        # Left side buttons
        ctk.CTkButton(
            actions_inner, text="⚙️  הגדרות", width=120, height=self._BTN_H,
            corner_radius=self._BTN_RADIUS, fg_color="#546E7A", hover_color="#455A64",
            font=("Calibri", 13), command=self._open_settings
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            actions_inner, text="📊  סטטיסטיקה", width=130, height=self._BTN_H,
            corner_radius=self._BTN_RADIUS, fg_color="#455A64", hover_color="#37474F",
            font=("Calibri", 13), command=lambda: StatsWindow(self)
        ).pack(side="left", padx=4)
        
        # Right side buttons
        ctk.CTkButton(
            actions_inner, text="📁  ייצוא לאקסל", width=140, height=self._BTN_H,
            corner_radius=self._BTN_RADIUS, fg_color=self._ACCENT_PURPLE, hover_color="#6D28D9",
            font=("Calibri", 13, "bold"), command=self._export_excel
        ).pack(side="right", padx=4)
        ctk.CTkButton(
            actions_inner, text="\u200Fשלח ב-WhatsApp 📱", width=165, height=self._BTN_H,
            corner_radius=self._BTN_RADIUS, fg_color=COLOR_WHATSAPP, hover_color="#1FAD55",
            font=("Calibri", 13, "bold"), command=self._send_whatsapp
        ).pack(side="right", padx=4)
        ctk.CTkButton(
            actions_inner, text="✓  אשר ושמור", width=140, height=self._BTN_H,
            corner_radius=self._BTN_RADIUS, fg_color=self._ACCENT_BLUE, hover_color="#1D4ED8",
            font=("Calibri", 13, "bold"), command=self._confirm_save
        ).pack(side="right", padx=4)
        ctk.CTkButton(
            actions_inner, text="▶  שבץ אוטומטי", width=150, height=self._BTN_H,
            corner_radius=self._BTN_RADIUS, fg_color=self._ACCENT_GREEN, hover_color="#047857",
            font=("Calibri", 13, "bold"), command=self._run_scheduler_clicked
        ).pack(side="right", padx=4)

        # ═══════════════════ Dashboard Footer ═══════════════════
        dash_card = ctk.CTkFrame(self, fg_color=self._CARD_BG, corner_radius=14, border_width=1, border_color=self._CARD_BORDER)
        dash_card.pack(fill="x", padx=20, pady=(0, 15))
        
        # Dashboard header stripe
        dash_header = ctk.CTkFrame(dash_card, fg_color="#1E293B", corner_radius=14, height=38)
        dash_header.pack(fill="x", padx=0, pady=0)
        dash_header.pack_propagate(False)
        
        ctk.CTkLabel(
            dash_header, text="📈  לוח בקרה והוגנות",
            font=("Calibri", 15, "bold"), text_color="white"
        ).pack(side="right", padx=20)
        
        # Scoreboard column headers
        self.scores_container = ctk.CTkFrame(dash_card, fg_color="transparent")
        self.scores_container.pack(fill="x", padx=0, pady=0)
        
        sb_header = ctk.CTkFrame(self.scores_container, fg_color="#F1F5F9", height=32, corner_radius=0)
        sb_header.pack(fill="x", padx=10, pady=(8, 0))
        sb_header.pack_propagate(False)
        
        headers_rtl = [
            ("שם העובד", 130),
            ("לילה חודש", 72),
            ("יום חודש", 72),
            ("לילה שבוע", 72),
            ("יום שבוע", 72),
            ("סטטוס", 88),
            ("מגמה", 160),
        ]
        for text, w in headers_rtl:
            ctk.CTkLabel(
                sb_header, text=text, width=w,
                font=("Calibri", 12, "bold"), text_color=self._TEXT_SECONDARY, anchor="center"
            ).pack(side="right", padx=10)

        self.scores_frame = ctk.CTkScrollableFrame(
            self.scores_container, height=220, fg_color="white", corner_radius=0
        )
        self.scores_frame.pack(fill="x", padx=10, pady=(0, 10))

    def _load_saved_url(self):
        url = self.settings.get("sheet_url", "")
        if url: self.url_entry.insert(0, url)

    # -------------------------------------------------------------- Networking
    def _import_clicked(self):
        url = self.url_entry.get().strip()
        if not url:
            messagebox.showerror("שגיאה", "יש להזין כתובת גיליון Google Sheets")
            return
        self.import_btn.configure(state="disabled")
        self._set_status("מייבא נתונים ומנתח זמינות...", "orange")
        threading.Thread(target=self._import_thread, args=(url,), daemon=True).start()

    def _import_thread(self, url):
        try:
            data = fetch_and_parse(url)
            self.after(0, lambda: self._import_done(url, data))
        except Exception as exc:
            self.after(0, lambda: self._import_error(exc))

    def _import_done(self, url, data):
        self.week_data = data
        self.settings["sheet_url"] = url
        save_settings(self.settings)
        self.current_assignments.clear()
        self.manual_flags.clear()
        self._build_schedule_table()
        self._set_status(f"הייבוא הושלם: {len(data['employees'])} עובדים זוהו", "green")
        self.import_btn.configure(state="normal")
        self._update_scores_display()

    def _import_error(self, exc):
        self._set_status("תקלה בייבוא הנתונים", "red")
        self.import_btn.configure(state="normal")
        messagebox.showerror("שגיאה", str(exc))

    # -------------------------------------------------------- Table Rendering
    def _build_schedule_table(self):
        """Construct the visual grid with Heatmap logic (כולל שני סלוטי בוקר בראשון–שלישי)."""
        for w in self.table_frame.winfo_children():
            w.destroy()
        self.schedule_menus.clear()
        data = self.week_data
        if data is None:
            return

        num_days = int(data["num_days"])
        employees = data["employees"]
        all_vals = [EMPTY_LABEL] + employees

        ctk.CTkLabel(self.table_frame, text="", width=150).grid(row=0, column=num_days)
        for d in range(num_days):
            col = num_days - 1 - d
            header = f"{data['day_names'][d]}\n{data['dates'][d]}"
            ctk.CTkLabel(
                self.table_frame, text=header, font=("Calibri", 11, "bold"),
                text_color="#1E293B", width=130,
            ).grid(row=0, column=col, padx=3, pady=10)

        for r_idx, (row_label, stype, slot) in enumerate(SCHEDULE_ROW_DEFS):
            row = r_idx + 1
            ctk.CTkLabel(
                self.table_frame, text=row_label, font=("Calibri", 11, "bold"),
                text_color="#37474F", anchor="e", width=150,
            ).grid(row=row, column=num_days, padx=10)
            for d in range(num_days):
                col = num_days - 1 - d
                valid_slots = slots_for_day_shift(d, stype)
                if slot not in valid_slots:
                    ctk.CTkLabel(
                        self.table_frame, text="—", text_color="#B0BEC5", width=125,
                    ).grid(row=row, column=col, padx=3, pady=5)
                    continue
                avail_dict = data.get("availability", {})
                avail_count: int = sum(
                    1 for e in employees if avail_dict.get(e, {}).get(d, {}).get(stype)
                )
                bg = COLOR_HEATMAP_EMPTY if avail_count == 0 else COLOR_HEATMAP_LOW if avail_count == 1 else "gray90"
                menu = ctk.CTkOptionMenu(
                    self.table_frame, values=all_vals, width=125, height=34,
                    fg_color=bg, button_color=bg, text_color="black", dynamic_resizing=False,
                    command=lambda val, dd=d, ss=stype, sl=slot: self._on_cell_change(dd, ss, sl, val),
                )
                menu.set(EMPTY_LABEL)
                menu.grid(row=row, column=col, padx=3, pady=5)
                self.schedule_menus[(d, stype, slot)] = menu
        self._update_cell_colors()

    def _on_cell_change(self, d, s, slot, v):
        key = (d, s, slot)
        if v == EMPTY_LABEL:
            self.current_assignments.pop(key, None)
            self.manual_flags.pop(key, None)
        else:
            self.current_assignments[key] = v
            self.manual_flags[key] = True
        self._update_cell_colors()
        self._update_scores_display()

    def _update_cell_colors(self):
        """Color based on assignment type or Heatmap if empty."""
        data = self.week_data
        if data is None:
            return
        for (d, s, slot), menu in self.schedule_menus.items():
            key = (d, s, slot)
            if key in self.current_assignments:
                emp = self.current_assignments[key]
                is_avail = data["availability"].get(emp, {}).get(d, {}).get(s, False)
                if self.manual_flags.get(key):
                    color = COLOR_MANUAL if is_avail else COLOR_WARNING
                    menu.configure(fg_color=color, button_color=color, text_color="white")
                else:
                    menu.configure(fg_color=COLOR_AUTO, button_color=COLOR_AUTO, text_color="white")
            else:
                avail_dict = data.get("availability", {})
                employees_list = data.get("employees", [])
                avail_count: int = sum(
                    1 for e in employees_list
                    if avail_dict.get(e, {}).get(d, {}).get(s)
                )
                bg = COLOR_HEATMAP_EMPTY if avail_count == 0 else COLOR_HEATMAP_LOW if avail_count == 1 else "gray90"
                menu.configure(fg_color=bg, button_color=bg, text_color="black")

    # ----------------------------------------------------------- Solver Logic
    def _run_scheduler_clicked(self):
        if not self.week_data: return
        self._set_status("מחשב את השיבוץ האופטימלי...", "orange")
        threading.Thread(target=self._scheduler_thread, daemon=True).start()

    def _scheduler_thread(self):
        try:
            data = self.week_data
            if data is None:
                self.after(0, lambda: self._set_status("אין נתונים לשיבוץ", "red"))
                return
            manual = {}
            for k, v in self.current_assignments.items():
                nk = normalize_assignment_key(k)
                if self.manual_flags.get(k) or self.manual_flags.get(nk):
                    manual[nk] = v
            history = get_current_month_scores()
            excluded = self.settings.get("excluded_employees", [])
            prev_shifts, prev_streaks = get_prev_week_data()

            result = schedule_shifts(
                availability=data["availability"],
                preferences=data.get("preferences", {}),
                day_names=data["day_names"],
                employees=data["employees"],
                manual_assignments=manual,
                historical_scores=history,
                excluded_employees=excluded,
                prev_week_shifts=prev_shifts,
                prev_work_streaks=prev_streaks,
                min_shifts_val=self.settings.get("min_shifts", 3),
                max_days_val=self.settings.get("max_consecutive_days", 6)
            )
            self.after(0, lambda r=result, m=manual: self._scheduler_done(r, m))
        except Exception as exc:
            self.after(0, lambda: self._scheduler_error(exc))

    def _scheduler_done(self, result, manual):
        if result.get("needs_night_approval"):
            detail = "\n".join(
                f"• {e}: {result.get('night_counts', {}).get(e, 0)} לילות"
                for e in (result.get("multi_night_employees") or [])
            )
            msg = (
                "הפתרון דורש שעובדים מסוימים יקבלו יותר מלילה אחד בשבוע "
                "(אין מספיק מתנדבי לילה או אילוצים אחרים).\n\n"
                f"{detail}\n\nלאשר את השיבוץ?"
            )
            if not messagebox.askyesno("אישור לילה כפול", msg, parent=self):
                self._set_status("השיבוץ בוטל — לא אושר לילה כפול", "orange")
                return

        self.current_assignments = dict(manual)
        self.manual_flags = {k: True for k in manual}
        for k, emp in result["assignments"].items():
            nk = normalize_assignment_key(k)
            if nk not in manual:
                self.current_assignments[nk] = emp
                self.manual_flags[nk] = False

        for key, menu in self.schedule_menus.items():
            emp = self.current_assignments.get(key)
            menu.set(emp if emp else EMPTY_LABEL)

        self._update_cell_colors()
        self._update_scores_display()
        self._set_status(
            "השיבוץ הושלם בהצלחה!" if not result["unfilled"] else "שיבוץ חלקי - יש משבצות ריקות",
            "green" if not result["unfilled"] else "orange",
        )

    def _scheduler_error(self, exc):
        self._set_status("שגיאה באלגוריתם", "red")
        messagebox.showerror("שגיאה", str(exc))

    # -------------------------------------------------------------- RTL FIXED Scoreboard
    def _update_scores_display(self):
        """לוח הוגנות: לילה (מספר) מול יום (בוקר+ערב+כונן משוקלל)."""
        for w in self.scores_frame.winfo_children():
            w.destroy()
        data = self.week_data
        if not data:
            return

        split_w = calculate_split_week_scores(self.current_assignments, data["day_names"])
        hist_split = get_current_month_split()
        excluded = self.settings.get("excluded_employees", [])
        all_emp = data["employees"]

        combined = {}
        for emp in all_emp:
            sw = split_w.get(emp, {"לילה": 0, "יום": 0.0})
            hm = hist_split.get(emp, {"day": 0.0, "night": 0})
            combined[emp] = float(sw["יום"]) + float(hm["day"]) + int(sw["לילה"]) + int(hm["night"])

        active_totals = [combined[e] for e in all_emp if e not in excluded]
        max_t = float(max(active_totals)) if active_totals else 1.0
        min_t = float(min(active_totals)) if active_totals else 0.0

        for r, emp in enumerate(all_emp):
            sw = split_w.get(emp, {"לילה": 0, "יום": 0.0})
            hm = hist_split.get(emp, {"day": 0.0, "night": 0})
            curr_t = float(combined[emp])
            is_exc = emp in excluded

            row_bg = "#F8FAFC" if r % 2 == 0 else "white"
            card = ctk.CTkFrame(self.scores_frame, fg_color=row_bg, corner_radius=6, height=34)
            card.pack(fill="x", padx=5, pady=1)

            if is_exc:
                tag, color = "מוחרג", "#9E9E9E"
            elif max_t > min_t:
                ratio = (curr_t - min_t) / (max_t - min_t)
                tag, color = (
                    ("גבוה ↑", "#E53935") if ratio > 0.7 else ("נמוך ↓", "#43A047") if ratio < 0.3 else ("מאוזן", "#42A5F5")
                )
            else:
                tag, color = "מאוזן", "#42A5F5"

            fill_w = max(int(160 * (curr_t / max_t if max_t > 0 else 0)), 2)

            ctk.CTkLabel(card, text=emp, font=("Calibri", 12, "bold"), text_color=self._TEXT_PRIMARY, width=130, anchor="e").pack(side="right", padx=6)
            ctk.CTkLabel(card, text=str(int(hm["night"])), font=("Calibri", 12), width=72, anchor="center").pack(side="right", padx=4)
            ctk.CTkLabel(card, text=f"{hm['day']:.1f}", font=("Calibri", 12), width=72, anchor="center").pack(side="right", padx=4)
            ctk.CTkLabel(card, text=str(int(sw["לילה"])), font=("Calibri", 12), width=72, anchor="center").pack(side="right", padx=4)
            ctk.CTkLabel(card, text=f"{sw['יום']:.1f}", font=("Calibri", 12), width=72, anchor="center").pack(side="right", padx=4)
            ctk.CTkLabel(card, text=tag, text_color=color, font=("Calibri", 11, "bold"), width=88, anchor="center").pack(side="right", padx=4)

            bar_cont = ctk.CTkFrame(card, width=160, height=12, fg_color="#E2E8F0", corner_radius=6)
            bar_cont.pack(side="right", padx=12)
            bar_cont.pack_propagate(False)
            ctk.CTkFrame(bar_cont, width=fill_w, height=12, fg_color=color, corner_radius=6).place(x=0, y=0)

    # -------------------------------------------------------- Messaging & Export
    def _send_whatsapp(self):
        data = self.week_data
        if data is None or not self.current_assignments: return
        assert data is not None
        msg = WHATSAPP_MSG_HEADER
        for d in range(data["num_days"]):
            msg += f"📅 *{data['day_names'][d]} ({data['dates'][d]}):*\n"
            for s in SHIFT_TYPES:
                slots = slots_for_day_shift(d, s)
                if len(slots) == 1:
                    emp = self.current_assignments.get((d, s, 0), "---")
                    msg += f"• {s}: {emp}\n"
                else:
                    for sl in slots:
                        emp = self.current_assignments.get((d, s, sl), "---")
                        msg += f"• {s} ({sl + 1}): {emp}\n"
            msg += "\n"
        webbrowser.open(f"https://web.whatsapp.com/send?text={urllib.parse.quote(msg)}")

    def _confirm_save(self):
        data = self.week_data
        if not data:
            return
        split = calculate_split_week_scores(self.current_assignments, data["day_names"])
        day_d = {}
        night_d = {}
        for emp in data["employees"]:
            s = split.get(emp, {"יום": 0.0, "לילה": 0})
            day_d[emp] = s["יום"]
            night_d[emp] = s["לילה"]
        update_month_scores_split(day_d, night_d)
        save_last_week_data(self.current_assignments, data["num_days"], data["employees"])
        messagebox.showinfo("נשמר", "הסידור והניקוד נשמרו בהצלחה!")
        self._update_scores_display()

    def _open_settings(self):
        data = self.week_data
        employees = data["employees"] if data else []
        excluded = self.settings.get("excluded_employees", [])
        dialog = SettingsDialog(self, employees, excluded, self.settings)
        self.wait_window(dialog)
        save_settings(self.settings)
        self._update_scores_display()
        if data: self._update_cell_colors()

    def _export_excel(self):
        data = self.week_data
        if not data: return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], title="שמור לוח משמרות", initialfile="schedule_noc.xlsx")
        if not path: return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "לוח משמרות"
            ws.sheet_view.rightToLeft = True
            
            # Formatting
            header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            center = Alignment(horizontal="center", vertical="center")
            border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

            ws.cell(1, 1, "משמרת").font = header_font
            ws.cell(1, 1).fill = header_fill
            ws.cell(1, 1).alignment = center
            ws.cell(1, 1).border = border
            ws.column_dimensions["A"].width = 20

            for d in range(data["num_days"]):
                cell = ws.cell(1, d + 2, f"{data['day_names'][d]} {data['dates'][d]}")
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = border
                ws.column_dimensions[cell.column_letter].width = 18

            excel_row = 2
            for row_label, stype, slot in SCHEDULE_ROW_DEFS:
                ws.cell(excel_row, 1, row_label).font = Font(bold=True)
                ws.cell(excel_row, 1).border = border
                for d in range(data["num_days"]):
                    emp = ""
                    if slot in slots_for_day_shift(d, stype):
                        emp = self.current_assignments.get((d, stype, slot), "")
                    cell = ws.cell(excel_row, d + 2, emp)
                    cell.alignment = center
                    cell.border = border
                    if emp:
                        key = (d, stype, slot)
                        color = "BBDEFB" if self.manual_flags.get(key) else "C8E6C9"
                        cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                excel_row += 1
            wb.save(path)
            messagebox.showinfo("ייצוא", "הקובץ נשמר בהצלחה!")
        except Exception as exc:
            messagebox.showerror("שגיאה", str(exc))

    def _set_status(self, text, color="gray"):
        self.status_label.configure(text=f"● {text}", text_color=color)

